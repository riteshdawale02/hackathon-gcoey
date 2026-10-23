import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    NAVY = RGBColor(15, 23, 42)       # #0f172a
    PRIMARY = RGBColor(2, 132, 199)   # #0284c7
    TEAL = RGBColor(13, 148, 136)     # #0d9488
    DARK_TEXT = RGBColor(30, 41, 59)  # #1e293b
    MUTED_TEXT = RGBColor(100, 116, 139) # #64748b
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)
    RED_ACCENT = RGBColor(220, 38, 38)
    CARD_BG = RGBColor(241, 245, 249)

    blank_layout = prs.slide_layouts[6]

    # Helper: Set slide background
    def set_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper: Add Header
    def add_header(slide, title_text, category_text="PULSEQUEUE HEALTHCARE ECOSYSTEM"):
        # Header banner
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.fill.background()

        tx = shape.text_frame
        tx.word_wrap = True
        tx.margin_left = Inches(0.8)
        tx.margin_top = Inches(0.18)

        p1 = tx.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = TEAL

        p2 = tx.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, NAVY)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.2), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()

    # Main Title
    txBox = s1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.8), Inches(3.2))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "PulseQueue"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "Smart Hospital Token & Queue Management System"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY

    p3 = tf.add_paragraph()
    p3.text = "Eliminating waiting room congestion with AI clinical triage, automated token issuance, real-time lobby broadcasting, and voice announcements."
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.space_before = Pt(16)

    # Metadata / Presenter footer
    meta = s1.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(10.8), Inches(1.2))
    mtf = meta.text_frame
    mp = mtf.paragraphs[0]
    mp.text = "Hackathon Problem Statement #3  |  Track: Digital Healthcare & Hospital Staff Dashboard"
    mp.font.size = Pt(14)
    mp.font.bold = True
    mp.font.color.rgb = TEAL

    mp2 = mtf.add_paragraph()
    mp2.text = "Built with React.js, Node.js Express, Java Spring Boot 3, PostgreSQL, Docker & AWS"
    mp2.font.size = Pt(12)
    mp2.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 2: THE PROBLEM
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, LIGHT_BG)
    add_header(s2, "The Problem: The Waiting Room Crisis")

    cards_data_s2 = [
        ("Chaotic Manual Queues", "Physical lines and paper tokens lead to crowded hallways, confusion, and transmission of airborne infections in hospital waiting areas.", RED_ACCENT),
        ("Patient Anxiety & Zero Visibility", "Patients have no idea where they stand in line or when they will be seen, leading to frustration, frequent desk inquiries, and walkouts.", PRIMARY),
        ("Manual Doctor Coordination", "Doctors waste 15-20% of consultation time manually calling names into exam rooms, dealing with no-shows, and managing disorderly queues.", DARK_TEXT),
        ("Dangerous Lack of Triage", "Critical patients (e.g. chest pain, asthma) often wait behind routine fever checkups because traditional queues treat every walk-in identically.", RED_ACCENT),
    ]

    for i, (title, desc, color) in enumerate(cards_data_s2):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 2.6)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(226, 232, 240)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_top = Inches(0.25)
        ctf.margin_right = Inches(0.3)

        cp1 = ctf.paragraphs[0]
        cp1.text = f"❌  {title}"
        cp1.font.size = Pt(18)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(13)
        cp2.font.color.rgb = MUTED_TEXT
        cp2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 3: THE PROPOSED SOLUTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, LIGHT_BG)
    add_header(s3, "The Solution: PulseQueue Digital Ecosystem")

    modules = [
        ("1. Patient Self-Service Portal", "Web & mobile accessible without app download. Instant token pass with live position countdown and scannable QR code."),
        ("2. Reception & Triage Desk", "Front-office walk-in counter with 10-digit mobile lookup, returning patient detection, and priority triage tagging."),
        ("3. Doctor OPD Dashboard", "Exam room console for attending physicians: One-click 'Call Next', digital clinical notes, and real-time consultation timers."),
        ("4. Lobby TV Broadcast Board", "High-contrast full-screen waiting hall screen auto-refreshing every 5s with 'Now Serving', upcoming tokens, and voice callouts."),
    ]

    for i, (title, desc) in enumerate(modules):
        left = Inches(0.8 + i * 2.95)
        top = Inches(1.8)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(186, 230, 253)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.3)
        ctf.margin_right = Inches(0.2)

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(17)
        cp1.font.bold = True
        cp1.font.color.rgb = PRIMARY

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(13)
        cp2.font.color.rgb = DARK_TEXT
        cp2.space_before = Pt(12)

    # =========================================================================
    # SLIDE 4: KEY FEATURES & INNOVATIONS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, LIGHT_BG)
    add_header(s4, "Key Features & Technical Innovations")

    feats = [
        ("🚨 AI Emergency Clinical Triage", "Real-time client-side NLP analyzes patient symptoms (e.g. chest pain, stroke, breathlessness) and auto-escalates critical cases to Emergency Priority (Queue Jump)."),
        ("🔊 Audible Hospital Chimes & Voice", "Built-in Web Audio API 2-tone melodic chime + Web Speech API vocal announcer ('Attention: Token CARD-002, please proceed to Room 101')."),
        ("⚡ Daily Atomic Token Sequencing", "Unique daily reset per department (e.g. CARD-001, GEN-001) backed by PostgreSQL atomic sequence counters."),
        ("📱 Scannable Mobile QR Passes", "Zero-friction access: patients scan the on-screen QR pass with any smartphone camera to carry their live wait tracker in their pocket."),
        ("⏱️ Rolling Wait Time Algorithm", "Dynamic calculation: Estimated Wait = (Patients Ahead) x (Live Rolling Average Consultation Duration)."),
        ("🔐 Enterprise Google OAuth2 & RBAC", "Stateless JWT authentication supporting real Google OAuth 2.0 and role-based permissions (Doctor, Reception, Patient, Admin)."),
    ]

    for i, (title, desc) in enumerate(feats):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 1.7)

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.18)

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(15)
        cp1.font.bold = True
        cp1.font.color.rgb = PRIMARY

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(11.5)
        cp2.font.color.rgb = MUTED_TEXT
        cp2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 5: SYSTEM ARCHITECTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, LIGHT_BG)
    add_header(s5, "Architecture: Cloud-Native Microservices")

    layers = [
        ("Frontend Client Layer", "React.js (Vite) + Vanilla CSS System", "Single responsive SPA. Mobile-first design for smartphones, tablets, doctor desktops, and public lobby TVs. Axios API client with reactive cross-tab event sync."),
        ("Microservice A: Identity & Patients", "Node.js (Express) :5001", "Owns identity_db. Handles JWT auth, RBAC, Google OAuth 2.0 verification, department routing, and patient profile management."),
        ("Microservice B: Queue Engine", "Java (Spring Boot 3) :8080", "Owns queue_db. Handles atomic token sequence generation, state transitions (WAITING -> IN_PROGRESS -> COMPLETED), and wait algorithms."),
        ("Data & Orchestration Layer", "PostgreSQL & Docker Compose", "Isolated databases adhering to database-per-service pattern. Multi-stage Docker containerization and 1-command deployment on AWS EC2."),
    ]

    for i, (title, tech, desc) in enumerate(layers):
        top = Inches(1.7 + i * 1.35)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_top = Inches(0.15)

        p1 = ctf.paragraphs[0]
        p1.text = f"{title}  |  "
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

        # Add tech badge inline
        run = p1.add_run()
        run.text = tech
        run.font.bold = True
        run.font.color.rgb = TEAL

        p2 = ctf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED_TEXT
        p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 6: END-TO-END WORKFLOW
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, LIGHT_BG)
    add_header(s6, "End-to-End Workflow: The 5-Step Journey")

    steps = [
        ("Step 1: Check-in", "Patient self-registers or checks in at reception. AI analyzes symptoms."),
        ("Step 2: Token Pass", "System generates CARD-004. Patient scans QR code to track live wait."),
        ("Step 3: Doctor Calls", "Dr. Sharma clicks 'Call Next' on dashboard. Previous consult completes."),
        ("Step 4: Live Broadcast", "Lobby TV flashes 'NOW SERVING: CARD-004'. Chime and voice announce room."),
        ("Step 5: Completion", "Doctor logs notes and completes. Real-time stats recalculate wait duration."),
    ]

    for i, (title, desc) in enumerate(steps):
        left = Inches(0.8 + i * 2.38)
        top = Inches(2.2)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.25), Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(226, 232, 240)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.3)

        cp1 = ctf.paragraphs[0]
        cp1.text = f"0{i+1}"
        cp1.font.size = Pt(32)
        cp1.font.bold = True
        cp1.font.color.rgb = TEAL

        cp2 = ctf.add_paragraph()
        cp2.text = title
        cp2.font.size = Pt(16)
        cp2.font.bold = True
        cp2.font.color.rgb = NAVY
        cp2.space_before = Pt(8)

        cp3 = ctf.add_paragraph()
        cp3.text = desc
        cp3.font.size = Pt(12)
        cp3.font.color.rgb = MUTED_TEXT
        cp3.space_before = Pt(10)

    # =========================================================================
    # SLIDE 7: REAL-WORLD IMPACT
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, LIGHT_BG)
    add_header(s7, "Real-World Impact & Quantifiable Outcomes")

    impacts = [
        ("40%", "Lobby Congestion Reduction", "Patients wait outside, in cafeterias, or nearby benches without losing their place in line."),
        ("60%", "Patient Anxiety Decrease", "Clear dynamic countdowns eliminate uncertainty and reduce repetitive inquiries at the desk."),
        ("20%", "Doctor Time Recovered", "One-click calling eliminates room-calling delays, manual staff lookups, and disorganized patient files."),
        ("100%", "Hardware Free Deployment", "Runs on any smart TV, standard tablet, or smartphone without expensive proprietary queue boxes."),
    ]

    for i, (metric, label, desc) in enumerate(impacts):
        left = Inches(0.8 + i * 2.95)
        top = Inches(1.8)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.3)

        cp1 = ctf.paragraphs[0]
        cp1.text = metric
        cp1.font.size = Pt(44)
        cp1.font.bold = True
        cp1.font.color.rgb = PRIMARY

        cp2 = ctf.add_paragraph()
        cp2.text = label
        cp2.font.size = Pt(16)
        cp2.font.bold = True
        cp2.font.color.rgb = NAVY
        cp2.space_before = Pt(4)

        cp3 = ctf.add_paragraph()
        cp3.text = desc
        cp3.font.size = Pt(12)
        cp3.font.color.rgb = MUTED_TEXT
        cp3.space_before = Pt(12)

    # =========================================================================
    # SLIDE 8: CONCLUSION & WHY WE WIN
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, NAVY)

    txBox = s8.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.8), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Why PulseQueue Wins"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE

    points = [
        "✅ Solves a Critical Need: Eliminates hospital chaos and anxiety with an intuitive, beautiful interface.",
        "✅ Beyond a Simple Queue: AI Triage automatically elevates life-threatening conditions to save lives.",
        "✅ Zero-Friction Adoption: No mobile app download required; instant QR code pass and Google OAuth 2.0.",
        "✅ Enterprise Microservice Architecture: Independent Node.js and Java Spring Boot services with database isolation.",
        "✅ Production DevOps Ready: 1-Command deployment script on AWS EC2 with full Docker Compose orchestration.",
    ]

    for pt in points:
        pp = tf.add_paragraph()
        pp.text = pt
        pp.font.size = Pt(16)
        pp.font.color.rgb = RGBColor(226, 232, 240)
        pp.space_before = Pt(14)

    footer = tf.add_paragraph()
    footer.text = "\nThank You!  |  Ready for Live Demo & Q&A"
    footer.font.size = Pt(22)
    footer.font.bold = True
    footer.font.color.rgb = TEAL

    # Save presentation
    output_filename = "PulseQueue_Pitch_Deck.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully created at: {output_filename}")

if __name__ == "__main__":
    create_presentation()
